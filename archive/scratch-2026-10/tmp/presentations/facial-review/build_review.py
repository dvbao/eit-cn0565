"""Build an editable research-review deck from the user's PowerPoint template.

This is a fallback for environments without the presentation artifact runtime.
All content shapes are editable PowerPoint text/tables, not slide screenshots.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "Powerpoint_template.pptx"
OUTPUT = ROOT / "output/presentation/Facial_Bioimpedance_Roadmap_Literature_Review.pptx"
HELPERS = ROOT / "tmp/presentations/week6-build/build_teaching_deck.py"
spec = importlib.util.spec_from_file_location("teaching_helpers", HELPERS)
H = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(H)


def slide(prs, title, *, title_size=25):
    s = H.add_content_slide(prs, title, len(prs.slides) + 1, title_size=title_size)
    return s


def panel(s, x, y, w, h, heading, body, *, color=H.MAROON, size=12.0):
    H.add_card(s, x, y, w, h, heading, body, accent=color, fill=H.LIGHT,
               title_size=13.8, body_size=size)


def source(s, url, label="Full paper"):
    sh = H.add_text(s, 0.48, 5.02, 8.55, 0.18, f"{label}: {url}", size=7.8,
                    color=H.MUTED, margin=0)
    for p in sh.text_frame.paragraphs:
        for run in p.runs:
            run.hyperlink.address = url


def note(s, vi):
    s.notes_slide.notes_text_frame.text = "VIETNAMESE SPEAKER NOTE\n" + vi


def add_paper_pair(prs, item):
    code, short, url = item["code"], item["short"], item["url"]
    s = slide(prs, f"{code}  {short}", title_size=23.5)
    panels = [
        ("1  What they did", item["what"]),
        ("2  Why difficult", item["challenge"]),
        ("3  Their response", item["tackle"]),
        ("4  Method", item["method"]),
        ("5  Setup", item["setup"]),
    ]
    panel(s, 0.45, 0.84, 4.40, 1.16, *panels[0])
    panel(s, 5.07, 0.84, 4.40, 1.16, *panels[1], color=H.BLUE)
    panel(s, 0.45, 2.10, 4.40, 1.16, *panels[2])
    panel(s, 5.07, 2.10, 4.40, 1.16, *panels[3], color=H.BLUE)
    panel(s, 0.45, 3.36, 9.02, 1.34, *panels[4], size=12.4)
    source(s, url)
    note(s, item["vi_method"])

    s = slide(prs, f"{code}  {short}", title_size=23.5)
    panel(s, 0.45, 0.84, 4.40, 1.28, "6  Finding / conclusion", item["result"], size=11.8)
    panel(s, 5.07, 0.84, 4.40, 1.28, "Evidence boundary", item["limit"], color=H.RED, size=11.8)
    panel(s, 0.45, 2.26, 4.40, 1.24, "7  Takeaway", item["takeaway"], size=12.0)
    panel(s, 5.07, 2.26, 4.40, 1.24, "8  Our project", item["project"], color=H.BLUE, size=12.0)
    panel(s, 0.45, 3.66, 9.02, 1.02, "Editor's reading", item["editor"], color=H.MAROON, size=12.0)
    source(s, url)
    note(s, item["vi_result"])


PAPERS = [
    dict(code="P1", short="Liu et al. 2025 — Facial Gesture BioZ",
         url="https://doi.org/10.1145/3745900.3746120",
         what="Recognize cheek puff/suck gestures from facial bioimpedance, not image reconstruction.",
         challenge="Obtain task information without camera or dense facial EMG.",
         tackle="Place three contacts on the face, derive two BioZ channels, classify temporal patterns.",
         method="20 Hz data; overlapping 20-sample windows (~1 s); Random Forest (100 trees); reported 80/20 split.",
         setup="4 participants. One electrode below mouth; one on each cheek. Left/right/omni puff and suck; null also mentioned.",
         result="Mean accuracy 0.93 and macro-F1 0.92 in Table 1; task-dependent traces are visible.",
         limit="n=4; anterior cheek electrodes; split grouping is not reported clearly; no across-session/subject proof; no muscle imaging.",
         takeaway="Direct precedent: a few BioZ channels can encode facial task patterns. The numbers are not transferable to our harder placement.",
         project="Closest baseline for Paper A. Replicate task definitions where useful, but test jaw/ear sites, reapplication and held-out sessions.",
         editor="Puff/suck can change oral air and geometry; classification is not proof of isolated muscle conductivity.",
         vi_method="Ba điện cực, hai kênh, bốn người. Đây là BioZ classification chứ không phải EIT tái tạo ảnh.",
         vi_result="93% là kết quả trên study rất nhỏ với vị trí má, không được quảng bá thành kết quả của vị trí tai/hàm."),
    dict(code="P2", short="Levit et al. 2024 — soft EMG + IPG",
         url="https://doi.org/10.1088/2057-1976/ad28cb",
         what="Simultaneously observe facial muscle-related biopotentials and vascular bioimpedance (IPG).",
         challenge="Facial artery and interwoven muscle responses are difficult to monitor concurrently without bulky contacts.",
         tackle="Screen-print flexible dry electrode arrays; use one for facial sEMG and another along superficial temporal artery.",
         method="ICA on 16-contact biopotential array; four-contact IPG tracks impedance baseline and cardiac pulse.",
         setup="4 healthy volunteers, ages 27–30; face and forearm comparisons; facial lip/teeth tasks and temporal artery region.",
         result="Muscle activation accompanied higher IPG baseline; artery narrowing is inferred by their model. Face pulse amplitude is less repeatable than forearm.",
         limit="No direct vessel-radius measurement in every trial; n=4; local temporal-artery IPG, not EIT or task classifier.",
         takeaway="Task–rest BioZ contains vascular and mechanical contributions. Baseline subtraction does not leave only muscle.",
         project="Add pulse/time controls; do not promise muscle-specific attribution from remote jaw/ear BioZ without validation.",
         editor="This paper supports our need to model blood vessels; it does not contradict Paper A's task-recognition aim.",
         vi_method="IPG là impedance plethysmography: nhìn biến thiên impedance do mạch đập/thể tích máu, khác EIT imaging.",
         vi_result="Tăng baseline Z đi cùng hoạt động cơ; tác giả suy luận mạch hẹp lại qua model, không chứng minh mọi thay đổi đều do mạch."),
    dict(code="P3", short="Statland et al. 2016 — facial EIM in FSHD",
         url="https://doi.org/10.1002/mus.25065",
         what="Assess EIM reliability/validity in FSHD, including direct orbicularis oris measurement.",
         challenge="Smaller facial muscles, hair, and operator placement can degrade measurement repeatability.",
         tackle="Use a smaller handheld EIM array on bilateral orbicularis oris and multiple limb/trunk sites.",
         method="Measure resistance/reactance/phase at 50 and 100 kHz; compare with severity/strength; test–retest ICC.",
         setup="35 clinically affected FSHD participants; 18 retested within three weeks. Facial contact is local, not remote.",
         result="Limb/trunk reliability generally high; facial right ICC <0.5, left 0.84–0.91. Reactance related to disease metrics.",
         limit="Disease=FSHD, not post-stroke dysarthria. No dynamic facial-task recognition, 3D EIT or per-muscle activation percentage.",
         takeaway="Facial EIM is precedent, but reliability is a major challenge even with contact directly over target muscle.",
         project="Predefine redonning/session repeatability and contact-quality checks before clinical claims.",
         editor="Do not borrow its clinical correlations as proof that our ear/jaw classifier measures weakness.",
         vi_method="FSHD là bệnh teo cơ khác stroke; nghiên cứu dùng probe đặt ngay lên cơ vòng môi.",
         vi_result="ICC bên phải kém cho thấy dữ liệu mặt nhạy với placement/operator, vì thế cần test tháo-gắn lại."),
    dict(code="P4", short="Kusche et al. 2024 — dual-channel EIM",
         url="https://doi.org/10.1109/JSEN.2024.3359284",
         what="Build a wearable two-channel impedance spectrometer for real-time muscle-contraction detection.",
         challenge="Distinguish changing muscle impedance while accounting for frequency and electrode-interface error.",
         tackle="Tetrapolar local electrode pairs on two antagonistic forearm muscles; characterize instrument with phantom.",
         method="Complex magnitude and phase across 20–230 kHz; 25 spectra/s/channel; R–C–R equivalent-circuit explanation.",
         setup="Four electrodes per local muscle channel; phantom calibration and initial human forearm contractions.",
         result="Phantom statistical error <1% magnitude and <0.4° phase; initial forearm tasks show time/frequency-dependent responses.",
         limit="Not a face model, not remote placement, not cohort-level recognition. Their optimum frequency is not ours.",
         takeaway="Intra-/extracellular/membrane ideas describe frequency response; electrode contact and direction also matter.",
         project="Use frequency/phase and bench uncertainty only if CN0565 measurement chain supports them; validate on face separately.",
         editor="A circuit fitted to forearm is an effective model, not a literal map of individual facial cells.",
         vi_method="R_e, R_i, C_m diễn tả đường ngoài tế bào, trong tế bào và màng; giá trị forearm không thể bê sang mặt.",
         vi_result="Đây là chứng cứ công cụ đo phổ và contraction ở tay; chỉ mượn phương pháp, không mượn kết luận về mặt."),
    dict(code="P5", short="Liu et al. 2024 — iMove",
         url="https://doi.org/10.1109/PerCom59722.2024.10494489",
         what="Recognize six upper-body fitness activities with wrist-to-wrist BioZ and wrist IMU.",
         challenge="Wearable motion tracking may miss body-state cues; BioZ alone is not necessarily strongest.",
         tackle="Fuse BioZ+IMU; contrastive training can improve a deployed IMU-only model.",
         method="Leave-one-person-out evaluation; compare BioZ-only, IMU-only, fusion and BioZ-assisted training.",
         setup="10 subjects across 5 days; electrodes at two wrists; left-wrist IMU; six activity classes.",
         result="Mean macro-F1: BioZ 75.36%, IMU 81.49%, fusion 89.57%, trained IMU-only 84.71%.",
         limit="Upper-body movement, not face; conductive path spans the body and does not isolate a particular muscle.",
         takeaway="A classifier can use distributed BioZ signals without reconstructing an image; modality ablations matter.",
         project="Paper A should test BioZ against video/motion baselines, not assume BioZ is always superior.",
         editor="Their cross-person split is a stronger evaluation idea than a random overlapping-window split.",
         vi_method="Tín hiệu giữa hai cổ tay phản ánh nhiều thay đổi cơ thể; không được gọi là model từng cơ tay.",
         vi_result="BioZ-only yếu hơn IMU ở study này; fusion tốt hơn. Cần baseline công bằng cho project mình."),
    dict(code="P6", short="Liu et al. 2024 — iEat",
         url="https://doi.org/10.1038/s41598-024-67765-5",
         what="Detect dining actions and food categories from wrist-to-wrist impedance.",
         challenge="Activity-associated impedance variations include changing external body–food conductive loops.",
         tackle="Treat those changes as useful features, not necessarily unwanted artifacts.",
         method="Single two-electrode impedance channel; user-independent neural model; leave-one-person/meal-out analyses.",
         setup="One electrode on each wrist; 10 volunteers, 40 meals, four intake activities and seven food types.",
         result="Mean macro-F1 86.4% for four activities, 64.2% for seven food types.",
         limit="Not facial muscle sensing and not EIT. The external food/utensil path is central to the method.",
         takeaway="Task recognition may work through geometry/contact pathways even when muscle tissue properties are not isolated.",
         project="Decide whether oral-air/jaw/contact cues are legitimate for Paper A; exclude them only if claiming intrinsic muscle signal.",
         editor="A high classifier score cannot by itself explain the physiological mechanism.",
         vi_method="iEat cố ý dùng conductive path qua tay, thức ăn, dụng cụ; rất khác đường dẫn qua cơ mặt.",
         vi_result="Bài này nhắc rằng nhận biết task và giải thích cơ chế là hai kết quả khác nhau."),
    dict(code="P7", short="Schumann et al. 2021 — facial EMG atlas",
         url="https://doi.org/10.1371/journal.pone.0254932",
         what="Map relative activation patterns for 29 voluntary facial tasks, not measure BioZ.",
         challenge="Facial muscles overlap; isolated voluntary activation is uncommon, making single-muscle labels unreliable.",
         tackle="Record high-channel bilateral sEMG and display normalized activity on 3D facial illustrations.",
         method="10 mimic muscle groups; 48-channel sEMG data; task-specific color normalization.",
         setup="30 healthy men; 29 tasks including six German vowels, unilateral smiles, lip and cheek movements.",
         result="Most tasks co-activate several sampled muscles; purse lips and one-sided smile have dominant measured groups.",
         limit="Buccinator not recorded; only men; surface EMG crosstalk; red is relative WITHIN each task, not cross-task strength.",
         takeaway="Rename 'single-muscle anchor' to 'dominant-pattern task'. Use Table 1, not anatomy intuition alone.",
         project="Ground truth for task design and validation EMG; not proof BioZ differentiates those muscles.",
         editor="Five vowels are not inherited from this atlas: it measured six German vowels; choose language/IPA deliberately.",
         vi_method="Atlas là sEMG, n=30 nam, 29 task. Màu đỏ là tương đối trong từng task, không so độ mạnh giữa hai task.",
         vi_result="Buccinator không được đo, nên không thể trích atlas để nói puff/suck chỉ kích hoạt buccinator."),
    dict(code="P8", short="Rüschenschmidt et al. 2022 — ear EMG",
         url="https://doi.org/10.3390/diagnostics12010121",
         what="Develop a method for measuring all nine intrinsic/extrinsic ear muscles during facial tasks.",
         challenge="Ear muscles are tiny and poorly mapped; facial tasks may activate them unexpectedly.",
         tackle="Cadaver anatomy guides electrode spots; compare needle and surface EMG, then test tasks/patients.",
         method="Video-documented facial tasks; surface and selected needle EMG; unilateral patient comparison.",
         setup="2 cadaver hemifaces; 12 healthy people; 7 patients with postparalytic facial synkinesis.",
         result="Multiple ear muscles activate during most facial tasks; greatest measured activation with smiling.",
         limit="This is EMG, not BioZ/EIT; condition is peripheral synkinesis, not post-stroke dysarthria.",
         takeaway="Ear electrodes are not a biologically silent site. Ear-only recognition may rely on local ear co-activation.",
         project="Record auricular EMG in a validation subset or frame result as task recognition, not remote cheek-muscle imaging.",
         editor="Ear placement remains promising for unobtrusiveness, but its mechanism needs a direct control.",
         vi_method="Study này chứng minh cơ tai có thể hoạt động khi cười; không chứng minh BioZ quanh tai đo được cơ má.",
         vi_result="Nếu classifier tại tai nhận biết smile, cần kiểm tra nó dùng cơ tai tại chỗ hay thông tin từ cơ mặt phía trước."),
]


def build():
    prs = Presentation(SOURCE)
    for idx in range(len(prs.slides) - 1, -1, -1):
        H.delete_slide(prs, idx)

    # Opening: preserve the source theme and branded slide layouts.
    s = prs.slides.add_slide(prs.slide_layouts[1])
    for sh in s.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 0:
            sh.text = "FACIAL BIOIMPEDANCE\nRESEARCH ROADMAP"
        elif sh.is_placeholder and sh.placeholder_format.idx == 13:
            sh.text = "Task recognition to 3D EIT | 30 September 2026"
    note(s, "Mục tiêu: tách paper đầu tiên về phân loại task khỏi mục tiêu dài hạn 3D EIT và đánh giá proof cần có.")

    s = slide(prs, "The research question is three different claims")
    panel(s, 0.55, 0.96, 2.76, 3.38, "A  Few-electrode BioZ", "Can permitted jaw/ear sites distinguish a preselected small task set across sessions and redonning?\n\nOutput: task label, not an image.", size=14.0)
    panel(s, 3.60, 0.96, 2.76, 3.38, "B  3D difference EIT", "Can wider 3D electrode coverage localize coarse left/right facial-region changes and reject jaw/contact confounds?\n\nOutput: coarse change map + uncertainty.", color=H.BLUE, size=13.8)
    panel(s, 6.65, 0.96, 2.76, 3.38, "C  Individual muscle deficit", "Can a validated map identify a weak muscle and quantify clinically meaningful impairment?\n\nOutput: calibrated biomarker. Highest risk.", color=H.RED, size=13.7)
    H.add_text(s, 0.60, 4.52, 8.85, 0.27, "More electrodes do not automatically yield 3D localization or a '% weak' muscle score.", size=13, bold=True, color=H.MAROON, align=PP_ALIGN.CENTER, margin=0)
    note(s, "A có precedent trực tiếp từ Liu nhưng layout khác. B cần 3D geometry, coverage và phantom. C còn đòi ground truth lâm sàng.")

    s = slide(prs, "What we measure: rest versus task")
    panel(s, 0.60, 1.05, 4.15, 2.35, "Rest state, v₀", "Inject current. Sense boundary voltages. Skin, bone, blood, muscle, air, fluid and electrode contact are ALL present.", size=14)
    panel(s, 5.15, 1.05, 4.15, 2.35, "Task state, v₁", "Repeat the SAME drive/sense protocol while the participant performs a task. Anatomy, geometry, contact and blood may all change.", color=H.BLUE, size=14)
    H.add_text(s, 1.0, 3.76, 8.0, 0.65, "Δv = v₁ − v₀   ≠   isolated muscle activation", size=23, bold=True, color=H.MAROON, align=PP_ALIGN.CENTER, margin=0)
    H.add_text(s, 0.64, 4.55, 8.75, 0.22, "Interpret a task-dependent boundary pattern; one voltage sign does not identify a tissue or muscle.", size=11.5, color=H.MUTED, align=PP_ALIGN.CENTER, margin=0)
    note(s, "Reference không phải mặt trống. Phép trừ chỉ loại phần không đổi, không tự loại hình học, electrode contact hay mạch máu.")

    s = slide(prs, "Prior work map: what has and has not been shown")
    H.add_table(s, 0.35, 0.91, 9.32, 3.85,
                ["Study", "Face?", "Task recognition?", "Image?", "Key boundary"],
                [["Liu 2025", "Yes", "Yes, cheek BioZ", "No", "n=4; anterior sites"],
                 ["Levit 2024", "Yes", "No; EMG + IPG", "No", "Vascular signal"],
                 ["Statland 2016", "Yes", "No; clinical EIM", "No", "Local OO probe"],
                 ["Schumann 2021", "Yes", "EMG task atlas", "No EIT", "Mixed activation"],
                 ["Ear EMG 2022", "Ear", "EMG task response", "No EIT", "Local ear signal"],
                 ["Kim 2019 / Piccin 2023", "Lower face / neck", "Airway tasks", "2D / 3D airway", "Not muscles"],
                 ["Our A → B → C", "Target", "To prove", "To prove", "Unobtrusive sites" ]],
                widths=[1.56, 0.89, 2.24, 1.31, 3.32], font_size=10.5)
    note(s, "Không có hàng nào hiện tại cùng lúc chứng minh unobtrusive remote placement, 3D cơ mặt và weakness percentage.")

    for item in PAPERS:
        add_paper_pair(prs, item)

    s = slide(prs, "Adjacent EIT precedent: face, neck, partial coverage")
    panel(s, 0.48, 0.84, 4.34, 1.23, "Visentin 2017 — facial EIT", "8 electrodes on one side of one face; exaggerated vowels/smile; deformation-inclusive proof of concept, not validated muscle map.", size=11.5)
    panel(s, 5.03, 0.84, 4.48, 1.23, "Kim et al. 2019 — lower-face airway", "16 electrodes guided by MRI/landmarks; 10 healthy participants; compare EIT with MRI during swallowing-induced airway change.", color=H.BLUE, size=11.5)
    panel(s, 0.48, 2.20, 4.34, 1.23, "Piccin et al. 2023 — 3D neck", "32 electrodes in two bands; CT-derived 3D neck model; target is airway, not expression muscles.", size=11.5)
    panel(s, 5.03, 2.20, 4.48, 1.23, "What transfers to us", "Use anatomy to target field coverage, test geometry and reconstruction. Their counts/placements do not prove our jaw/ear layout.", color=H.BLUE, size=11.5)
    H.add_text(s, 0.47, 3.67, 8.96, 0.94,
               "FULL LINKS\nhttps://tsukuba.repo.nii.ac.jp/record/43368/files/DA08090.pdf\nhttps://doi.org/10.5664/jcsm.7714  |  https://doi.org/10.3389/frsle.2023.1238508",
               size=10.5, color=H.MUTED, margin=0.03)
    note(s, "Cả ba giúp biện luận placement, nhưng target là deformation/airway, không xác nhận cơ mặt yếu.")

    s = slide(prs, "Other facial BioZ / tissue-model precedents")
    panel(s, 0.52, 0.84, 4.22, 1.52, "Painometry 2020", "Local corrugator-region impedance plus EEG/PPG/GSR for pain inference. Not remote facial-task recognition.", size=12.2)
    panel(s, 5.08, 0.84, 4.40, 1.52, "Schooling et al. 2020/21", "Layered anisotropic tongue FEM + multifrequency Cole-type interpretation (41 ALS, 30 controls). Not facial expression.", color=H.BLUE, size=12.0)
    panel(s, 0.52, 2.55, 8.96, 1.04, "A 2026 update", "Büchner et al. use sEMG + 3D video (36 healthy adults) to visualize activation on a 3D face. That is NOT 3D EIT tomography.", size=12.3)
    H.add_text(s, 0.53, 3.79, 8.95, 0.83,
               "FULL LINKS\nhttps://doi.org/10.1145/3386901.3389022  |  https://doi.org/10.1088/1361-6579/abcb9b\nhttps://doi.org/10.1016/j.heliyon.2026.e45408",
               size=10.3, color=H.MUTED, margin=0.03)
    note(s, "Có model lưỡi và impedance cục bộ ở trán, nhưng không có model đã kiểm chứng cho toàn bộ mặt theo layout mình.")

    s = slide(prs, "Beginner's tissue model: where can Δv come from?")
    H.add_table(s, 0.35, 0.88, 9.30, 3.72,
                ["Component", "Physical effect", "Why it matters"],
                [["Electrode + skin", "Contact pressure/hydration", "Can mimic task or drift"],
                 ["Fat / connective tissue", "Depth and current path", "Attenuates/redirects ROI sensitivity"],
                 ["Muscle cells", "Extra-/intracellular + membrane", "Frequency/anisotropy; contraction geometry"],
                 ["Blood vessels", "Volume, pulse, constriction", "Task-linked vascular BioZ (Levit)"],
                 ["Bone / teeth", "Conductivity + geometry", "Rest field and moving jaw"],
                 ["Mouth/nose air + saliva", "Changing cavities/fluid paths", "Strong cue in puff/suck/speech"],
                 ["Ear muscles", "Local co-activation", "Ear site is not physiologically silent"]],
                widths=[1.75, 3.00, 4.55], font_size=11)
    H.add_text(s, 0.49, 4.69, 9.02, 0.25, "Each term can affect the voltage vector; anatomy at rest does not vanish after subtraction.", size=11.5, bold=True, color=H.MAROON, margin=0)
    note(s, "Nói đơn giản: current đi qua mọi tissue, nhưng sensor chỉ đọc điện áp ở biên. Task có thể thay nhiều thứ cùng lúc.")

    s = slide(prs, "From equivalent circuit to 3D forward model")
    panel(s, 0.53, 0.90, 4.24, 2.79, "Material model (one effective tissue)", "Rₑ: extracellular path\nRᵢ: intracellular path\nCₘ: membrane effect\n\nFrequency changes which pathways contribute. Cole model summarizes a broad spectrum; fit parameters are not a named facial muscle.", size=13.1)
    panel(s, 5.08, 0.90, 4.42, 2.79, "Anatomical + electrode model", "A 3D mesh assigns frequency-dependent conductivity/permittivity to skin, fat, muscle, bone, blood and cavities; a complete-electrode model handles contact.\n\nForward solve predicts boundary voltage for a chosen current pattern.", color=H.BLUE, size=13)
    H.add_text(s, 0.54, 3.95, 8.92, 0.74,
               "Material prior ≠ measured facial value. Use uncertainty ranges, anatomy registration, contact calibration and mesh-convergence checks.",
               size=13.1, bold=True, color=H.MAROON, align=PP_ALIGN.CENTER, margin=0)
    H.add_text(s, 0.52, 4.93, 8.9, 0.18, "https://doi.org/10.1109/JSEN.2024.3359284  |  https://doi.org/10.1088/1361-6579/abcb9b  |  https://itis.swiss/virtual-population/tissue-properties/database/low-frequency-conductivity", size=7.4, color=H.MUTED, margin=0)
    note(s, "Equivalent circuit giải thích phổ của một vùng; FEM 3D giải quyết dòng đi qua toàn anatomy và electrodes.")

    s = slide(prs, "Placement depends on the question being asked")
    H.add_table(s, 0.37, 0.89, 9.26, 3.67,
                ["Layout family", "Paper A: classifier", "Paper B: 3D image", "Specific control"],
                [["Ear only", "Wearable; test lower-face signal", "Likely weak coverage; must test", "Local ear EMG + head pose"],
                 ["Ear + jaw + submental", "Candidate compromise", "Maybe coarse lower-face ROI", "Jaw/air/contact ablations"],
                 ["+ temple/forehead", "More spatial diversity vs visibility", "Potential multi-height coverage", "Equal-count redistribution"],
                 ["Anterior cheek reference", "Temporary Liu-like comparator", "Improves target coverage; violates final constraint", "Same device/protocol where possible"]],
                widths=[1.69, 2.30, 2.67, 2.60], font_size=10.7)
    H.add_text(s, 0.55, 4.63, 8.9, 0.31, "Closed ring is not a theorem. Partial coverage is possible but inverse localization must be measured.", size=12.5, bold=True, color=H.MAROON, align=PP_ALIGN.CENTER, margin=0)
    source(s, "https://doi.org/10.5664/jcsm.7714", "Anatomy-informed example")
    note(s, "So layout theo cùng ngân sách điện cực. Classification và reconstruction có hàm mục tiêu khác nhau.")

    s = slide(prs, "The task list: correct the muscle labels")
    H.add_table(s, 0.37, 0.88, 9.25, 3.77,
                ["Proposed label", "Atlas Table 1 says", "Editorial correction"],
                [["Purse lips (T14)", "Orbicularis oris highest sampled", "Dominant pattern, not isolated"],
                 ["Unilateral smiles (T22/23)", "Zygomatic highest sampled", "Separate left/right + video"],
                 ["Lower lip protrusion (T11)", "Mentalis + depressors + OO", "Mixed, not single target"],
                 ["Lips wide, jaw shut (T20)", "Inferior OO + depressor labii + mentalis", "Not depressor anguli highest"],
                 ["Upper-lip lift (T12)", "Levators + OO", "Mixed"],
                 ["Lip press / corners down", "T7 / T8: multiple muscles", "Mixed activation"],
                 ["Cheek puff / suck (T15/16)", "OO + mentalis + depressor labii", "Buccinator was NOT sampled"]],
                widths=[2.25, 3.47, 3.53], font_size=10.6)
    source(s, "https://doi.org/10.1371/journal.pone.0254932", "Source: Schumann et al., Table 1")
    note(s, "OO là orbicularis oris. Highest chỉ trong các cơ được ghi và từng task, không phải cơ duy nhất.")

    s = slide(prs, "Vowels, jaw and cheek: what are we actually testing?")
    panel(s, 0.48, 0.83, 4.33, 1.25, "Vowels", "Atlas has SIX German vowels, not a universal five-vowel battery. Specify language, IPA and sustained vs speech production.", size=11.9)
    panel(s, 5.05, 0.83, 4.42, 1.25, "Jaw tasks", "Jaw-open/closed-lips is atlas T18. Wide jaw opening and clench need separate rationale; clench is a jaw control, not a facial-muscle anchor.", color=H.BLUE, size=11.8)
    panel(s, 0.48, 2.22, 4.33, 1.25, "Cheek puff/suck", "Useful bridge to Liu, but air volume, oral pressure, lip seal and geometry can all change BioZ. Atlas did not measure buccinator.", size=11.8)
    panel(s, 5.05, 2.22, 4.42, 1.25, "Clinical relevance", "Non-speech mimic tasks establish sensing feasibility; they do not by themselves prove dysarthria assessment or speech intelligibility.", color=H.RED, size=11.8)
    H.add_text(s, 0.48, 3.73, 8.95, 0.76,
               "Full links: https://doi.org/10.1371/journal.pone.0254932  |  https://doi.org/10.1177/00220345770560071301\nhttps://www.asha.org/Practice-Portal/Clinical-Topics/Dysarthria-in-Adults/",
               size=9.5, color=H.MUTED, margin=0.02)
    note(s, "Chưa nên đưa tất cả vào primary classes; chọn một core set nhỏ theo use case trước khi xem dữ liệu.")

    s = slide(prs, "Your 5 × 5 s protocol: calculated burden")
    panel(s, 0.50, 0.89, 4.28, 2.74, "If all 19 labels are used", "6 anchor labels + 3 mixed + 5 vowels + 3 jaw + 2 cheek = 19.\n\n5 repetitions each = 95 trials.\n5 s active each = 475 s active time.", size=14)
    panel(s, 5.07, 0.89, 4.38, 2.74, "Timing depends on rest convention", "With 5 s rest after every repetition: 475 s rest + 18 × 10 s between tasks = 1,130 s = 18 min 50 s.\n\nWithout final 5 s rest per task: 17 min 15 s.", color=H.BLUE, size=13.8)
    H.add_text(s, 0.55, 3.89, 8.90, 0.76, "Pilot timing, fatigue, signal recovery and task quality. These numbers are a proposal—not literature-validated BioZ parameters.", size=13.6, bold=True, color=H.MAROON, align=PP_ALIGN.CENTER, margin=0)
    note(s, "Chưa tính practice, hướng dẫn, tháo-gắn, lỗi trial. Hold 5 giây vowel là sustained phonation chứ không phải natural speech.")

    s = slide(prs, "Recommended study design: prespecify before classifying")
    panel(s, 0.45, 0.85, 4.40, 1.13, "Primary classes", "Rest, purse lips, left smile, right smile; optional lip press or cheek puff only with a clear use case.", size=11.8)
    panel(s, 5.05, 0.85, 4.40, 1.13, "Exploratory block", "Vowels, jaw and cheek tasks test generalization/confounds; do not select classes after test accuracy.", color=H.BLUE, size=11.8)
    panel(s, 0.45, 2.08, 4.40, 1.13, "Data and controls", "Sync BioZ with video; EMG subset (including ear if needed); record contact, pose, jaw and pulse.", size=11.8)
    panel(s, 5.05, 2.08, 4.40, 1.13, "Evaluation", "Split trials before windowing; report within-session, cross-session/redonning and unseen-person macro-F1 + uncertainty.", color=H.BLUE, size=11.8)
    panel(s, 0.45, 3.33, 9.00, 1.30, "A valid first-paper claim", "Constrained jaw/ear placement detects a preselected set of task patterns reliably under explicit nuisance controls. Do not claim individual-muscle imaging or post-stroke diagnosis.", size=12.2)
    note(s, "Paper A quan trọng nhất là reliability qua buổi, tháo-gắn và người mới. Video/EMG giúp kiểm tra task thực sự thực hiện.")

    s = slide(prs, "Proof ledger: what must the next experiment decide?")
    H.add_table(s, 0.34, 0.85, 9.31, 3.92,
                ["Claim", "Present evidence", "Required proof", "Status"],
                [["Task changes facial BioZ", "Liu, Levit", "Replicate at permitted sites", "Supported in general"],
                 ["Ear/jaw recognizes tasks", "No exact-layout study", "Matched tasks; held-out sessions", "Hypothesis"],
                 ["A task isolates one muscle", "Atlas shows co-activation", "EMG if needed", "False as stated"],
                 ["Δv is pure muscle", "Blood/air/contact pathways", "Nuisance controls + model", "False as stated"],
                 ["3D localizes facial ROIs", "Airway EIT precedent", "3D FEM + ground-truth phantom", "Unproven"],
                 ["Weakness percentage", "No direct evidence", "Defined scale + clinical validation", "Future hypothesis"]],
                widths=[2.13, 2.10, 3.00, 2.08], font_size=10.4)
    note(s, "Bảng này là cách editor nhìn paper: claim nào đã có prior, claim nào cần experiment, claim nào phải bỏ.")

    s = slide(prs, "Parameter provenance: every number needs a reason")
    H.add_table(s, 0.35, 0.86, 9.30, 3.90,
                ["Parameter", "Meaning", "How to defend value"],
                [["Electrode count / site", "Coverage and burden", "Equal-budget layout ablation + FEM"],
                 ["Drive/sense pattern", "Independent boundary views", "Rank/sensitivity + bench validation"],
                 ["Frequency / phase", "Tissue dispersion and instrument", "Verified hardware sweep; held-out ablation"],
                 ["Tissue σ, ε, anisotropy", "Material properties", "IT'IS/Gabriel range, uncertainty sweep"],
                 ["Contact + geometry", "Boundary condition", "Measure/reapply/scan + nuisance test"],
                 ["Trial timing / effort", "Physiological repeatability", "Pilot plateau, fatigue and recovery"],
                 ["Mesh h / regularization", "Numerical/inverse assumptions", "Convergence, phantom error, uncertainty"]],
                widths=[2.12, 2.58, 4.60], font_size=10.7)
    H.add_text(s, 0.41, 4.83, 9.16, 0.18, "Prior: https://itis.swiss/virtual-population/tissue-properties/database/low-frequency-conductivity", size=8.2, color=H.MUTED, margin=0)
    note(s, "Mỗi con số cần ghi paper, tissue, frequency, uncertainty, phiên bản, và phép sweep. Không có nguồn thì ghi unknown.")

    s = slide(prs, "Publication decision: staged, not promised")
    panel(s, 0.55, 0.93, 4.22, 3.33, "Paper A: feasible if signal survives", "A few unobtrusive electrodes + selected task recognition + cross-session/redonning proof + clear mechanism limits.\n\nNovelty is constrained placement and robust evaluation, not simply fewer classes.", size=14.1)
    panel(s, 5.10, 0.93, 4.35, 3.33, "Paper B: conditional and harder", "Multi-plane/wider 3D EIT + coarse ROI reconstruction + identifiability + phantom truth.\n\nOnly then consider single-muscle and clinical weakness calibration as a separate claim.", color=H.BLUE, size=14)
    H.add_text(s, 0.65, 4.45, 8.74, 0.43, "First actionable gate: bench/phantom system verification and ethics/safety review before any human current injection.", size=12.4, bold=True, color=H.RED, align=PP_ALIGN.CENTER, margin=0)
    note(s, "Không thể hứa chắc paper thứ hai. Nếu FEM cho thấy không đủ sensitivity ở layout cho phép, cần đổi placement hoặc đổi claim.")

    closing = prs.slides.add_slide(prs.slide_layouts[4])
    closing.shapes.title.text = "Questions / decisions"
    note(closing, "Cần chốt primary classes, ngôn ngữ nguyên âm, điều kiện anterior comparator và giới hạn unobtrusive cụ thể.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
