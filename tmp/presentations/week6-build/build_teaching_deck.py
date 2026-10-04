from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


MAROON = RGBColor(113, 31, 59)
MAROON_DARK = RGBColor(83, 24, 45)
MAROON_LIGHT = RGBColor(245, 235, 239)
RED = RGBColor(197, 27, 58)
BLUE = RGBColor(39, 111, 191)
NAVY = RGBColor(49, 72, 94)
TEXT = RGBColor(34, 34, 34)
MUTED = RGBColor(92, 92, 92)
LIGHT = RGBColor(244, 246, 248)
MID = RGBColor(222, 227, 231)
WHITE = RGBColor(255, 255, 255)


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def remove_body_placeholders(slide) -> None:
    for shape in list(slide.shapes):
        if not shape.is_placeholder:
            continue
        if shape.placeholder_format.type in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
            PP_PLACEHOLDER.SLIDE_NUMBER,
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
        ):
            slide.shapes._spTree.remove(shape._element)


def style_title(slide, title: str, size: float = 26) -> None:
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.left = Inches(0.40)
    title_shape.top = Inches(0.06)
    title_shape.width = Inches(9.05)
    title_shape.height = Inches(0.55)
    tf = title_shape.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = MAROON


def add_page_number(slide, number: int) -> None:
    add_text(
        slide,
        9.18,
        5.20,
        0.42,
        0.20,
        str(number),
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 15,
    bold: bool = False,
    color: RGBColor = TEXT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    rounded: bool = False,
    margin: float = 0.06,
    font: str = "Arial",
):
    if fill is None and line is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line
            shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_paragraph_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paragraphs: list[dict],
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    rounded: bool = False,
    margin: float = 0.10,
    valign=MSO_ANCHOR.TOP,
):
    shape = add_text(
        slide,
        x,
        y,
        w,
        h,
        "",
        fill=fill,
        line=line,
        rounded=rounded,
        margin=margin,
        valign=valign,
    )
    tf = shape.text_frame
    tf.clear()
    for idx, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(item.get("after", 5))
        p.line_spacing = item.get("line_spacing", 1.0)
        run = p.add_run()
        run.text = item["text"]
        run.font.name = item.get("font", "Arial")
        run.font.size = Pt(item.get("size", 14))
        run.font.bold = item.get("bold", False)
        run.font.italic = item.get("italic", False)
        run.font.color.rgb = item.get("color", TEXT)
    return shape


def add_card(slide, x, y, w, h, title, body, *, accent=MAROON, fill=LIGHT, title_size=17, body_size=12.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = MID
    shape.line.width = Pt(1)
    add_text(slide, x + 0.12, y + 0.10, w - 0.24, 0.30, title, size=title_size, bold=True, color=accent, margin=0)
    add_text(slide, x + 0.12, y + 0.48, w - 0.24, h - 0.58, body, size=body_size, color=TEXT, margin=0)
    return shape


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw = iw * scale
    ph = ih * scale
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - pw) / 2),
        Inches(y + (h - ph) / 2),
        width=Inches(pw),
        height=Inches(ph),
    )


def add_arrow(slide, x, y, w=0.42, h=0.28, color=MAROON):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_table(slide, x, y, w, h, headers, rows, widths=None, font_size=11.5):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    for c, value in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = value
        cell.fill.solid()
        cell.fill.fore_color.rgb = MAROON
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(font_size)
                r.font.bold = True
                r.font.color.rgb = WHITE
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 else RGBColor(247, 241, 243)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = TEXT
    return table_shape


def add_source(slide, text: str) -> None:
    add_text(slide, 0.55, 4.92, 8.20, 0.22, text, size=8.2, color=MUTED, margin=0)


def add_notes(slide, english: str, vietnamese: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"ENGLISH SCRIPT\n{english}\n\nVIETNAMESE EXPLANATION\n{vietnamese}"


def add_content_slide(prs: Presentation, title: str, number: int, *, layout_index=3, title_size=26):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    remove_body_placeholders(slide)
    style_title(slide, title, title_size)
    add_page_number(slide, number)
    return slide


def add_state_circle(slide, x, y, label, object_present: bool, accent):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(1.55), Inches(1.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(235, 239, 242)
    circle.line.color.rgb = NAVY
    circle.line.width = Pt(1.5)
    if object_present:
        target = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.85), Inches(y + 0.38), Inches(0.30), Inches(0.30))
        target.fill.solid()
        target.fill.fore_color.rgb = RED
        target.line.fill.background()
    for i in range(16):
        import math

        angle = math.pi / 2 - i * 2 * math.pi / 16
        ex = x + 0.775 + 0.77 * math.cos(angle) - 0.035
        ey = y + 0.775 - 0.77 * math.sin(angle) - 0.035
        electrode = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ex), Inches(ey), Inches(0.07), Inches(0.07))
        electrode.fill.solid()
        electrode.fill.fore_color.rgb = accent
        electrode.line.fill.background()
    add_text(slide, x, y + 1.70, 1.55, 0.30, label, size=16, bold=True, color=accent, align=PP_ALIGN.CENTER, margin=0)


def build(source: Path, assets: Path, output: Path) -> None:
    prs = Presentation(source)

    # Preserve slides 1-17. Remove the old simulation section and closing
    # slide; a fresh closing slide is added from the same branded layout.
    # Removing it avoids part-name collisions when python-pptx appends enough
    # new slides to reach the source deck's original slide28.xml.
    for index in range(27, 16, -1):
        delete_slide(prs, index)

    # 18 — section divider.
    slide = add_content_slide(prs, "3. Active Sensing - EIT Simulation Study", 18, title_size=25)
    add_text(slide, 0.85, 1.38, 8.25, 0.72, "From boundary measurements to a reconstructed conductivity-change image", size=24, bold=True, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, 1.30, 2.35, 7.40, 0.42, "Physical setup  →  208 voltage differences  →  BP / JAC / GREIT", size=18, color=MAROON, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, 1.00, 3.20, 8.00, 0.72, "Goal: understand what is simulated, what each parameter means, and how to compare algorithms fairly before testing hardware.", size=16, color=MUTED, align=PP_ALIGN.CENTER, margin=0)
    add_notes(slide,
        "This section starts from the physical measurement, not from the algorithms. I will first show where the 208 values come from, then explain the conductivity model, the finite-element mesh, the voltage data, and finally the three reconstruction methods.",
        "Phần này bắt đầu từ phép đo vật lý chứ chưa vào thuật toán ngay. Trình tự là: 16 điện cực tạo ra 208 hiệu điện thế, xây dựng mô hình conductivity, forward solve, rồi mới so sánh BP, JAC và GREIT.")

    # 19 — physical story.
    slide = add_content_slide(prs, "What EIT Actually Measures", 19)
    add_picture_contain(slide, assets / "eit_setup.png", 0.35, 0.70, 5.35, 4.05)
    add_card(slide, 5.85, 0.82, 3.55, 0.98, "1. Conductive medium", "Saline or tissue provides background conductivity σbg.", accent=MAROON, fill=LIGHT, title_size=15.5, body_size=11.5)
    add_card(slide, 5.85, 1.91, 3.55, 0.98, "2. Boundary current", "One adjacent pair injects +I and removes −I.", accent=MAROON, fill=LIGHT, title_size=15.5, body_size=11.5)
    add_card(slide, 5.85, 3.00, 3.55, 1.10, "3. Boundary voltage", "A conductivity change perturbs current paths and the voltage pattern measured by the sense pairs.", accent=MAROON, fill=LIGHT, title_size=15.5, body_size=11.2)
    add_text(slide, 5.88, 4.22, 3.50, 0.48, "EIT infers the object from the complete boundary pattern.", size=12.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=RGBColor(236, 243, 249), line=MID, rounded=True)
    add_source(slide, "Concept supported by Liu et al. (2018): current is injected through a boundary pair and voltage differences are measured on remaining pairs.")
    add_notes(slide,
        "A useful sentence is: EIT does not photograph the object. It applies current on the boundary and observes how the boundary-voltage pattern changes. A more conductive or less conductive object perturbs the current distribution. Reconstruction is the inverse step that estimates where that internal conductivity change occurred.",
        "Câu dễ nhớ: EIT không chụp trực tiếp vật thể. Nó bơm dòng ở biên, đo hiệu điện thế ở biên, rồi suy ngược vị trí thay đổi conductivity bên trong. Object chỉ được nhận ra vì nó làm thay đổi pattern điện áp.")

    # 20 — 208 measurements.
    slide = add_content_slide(prs, "Where the 208 Voltage Differences Come From", 20, title_size=25)
    add_card(slide, 0.55, 0.92, 2.30, 1.20, "16 drive pairs", "The adjacent drive rotates through all electrodes.", accent=RED, fill=RGBColor(255, 241, 244), title_size=18, body_size=11.3)
    add_arrow(slide, 2.95, 1.36)
    add_card(slide, 3.50, 0.92, 2.30, 1.20, "13 sense pairs", "Each drive retains 13 valid adjacent sense pairs.", accent=BLUE, fill=RGBColor(239, 247, 255), title_size=18, body_size=11.3)
    add_arrow(slide, 5.90, 1.36)
    add_card(slide, 6.45, 0.92, 2.65, 1.20, "208 / frame", "16 × 13 sequential voltage differences.", accent=MAROON, fill=MAROON_LIGHT, title_size=18, body_size=11.3)
    add_text(slide, 0.70, 2.46, 8.60, 0.34, "Example for drive pair E1-E2", size=18, bold=True, color=MAROON, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, 0.75, 2.96, 1.55, 0.52, "Invalid pairs", size=14, bold=True, color=RED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=RGBColor(255, 236, 240), line=RED, rounded=True)
    add_text(slide, 2.42, 2.96, 2.25, 0.52, "E16-E1, E1-E2, E2-E3", size=13.5, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=WHITE, line=MID, rounded=True)
    add_text(slide, 4.88, 2.96, 1.55, 0.52, "Valid pairs", size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=RGBColor(237, 246, 255), line=BLUE, rounded=True)
    add_text(slide, 6.55, 2.96, 2.65, 0.52, "E3-E4, E4-E5, ... , E15-E16", size=13.5, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=WHITE, line=MID, rounded=True)
    add_paragraph_box(slide, 0.80, 3.72, 8.40, 0.96, [
        {"text": "Why 13, not 12?", "size": 14.5, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 2},
        {"text": "Remove the two drive electrodes: 14 electrodes remain on one open arc, and 14 points form 13 adjacent sensing gaps.", "size": 12.2, "align": PP_ALIGN.CENTER},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.10)
    add_notes(slide,
        "A frame is not one simultaneous reading. It is a sequence. First drive E1-E2 and measure thirteen valid adjacent voltage pairs. Then rotate the drive to E2-E3 and repeat. Sixteen drive positions times thirteen values gives 208. The three adjacent pairs touching E1 or E2 are excluded in the E1-E2 example.",
        "Một frame là chuỗi phép đo. Với drive E1-E2, ba cặp sense chạm vào điện cực drive bị loại, còn 13 cặp hợp lệ. Lặp cho 16 vị trí drive nên được 16 × 13 = 208. Không phải 12 vì 14 điện cực còn lại tạo 13 khoảng liền kề trên cung mở.")

    # 21 — v0, v1, dv.
    slide = add_content_slide(prs, "Reference, Target, and Change Frames", 21)
    add_state_circle(slide, 0.72, 1.15, "Reference state  v₀", False, NAVY)
    add_arrow(slide, 2.55, 1.76, 0.50, 0.34)
    add_state_circle(slide, 3.25, 1.15, "Target state  v₁", True, RED)
    add_arrow(slide, 5.08, 1.76, 0.50, 0.34)
    add_paragraph_box(slide, 5.85, 1.05, 3.30, 2.12, [
        {"text": "Δv = v₁ − v₀", "size": 24, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 8},
        {"text": "208 measured changes", "size": 17, "bold": True, "align": PP_ALIGN.CENTER, "after": 7},
        {"text": "This change vector is the input to BP, JAC, and GREIT.", "size": 14, "align": PP_ALIGN.CENTER},
    ], fill=MAROON_LIGHT, line=MAROON, rounded=True, margin=0.13, valign=MSO_ANCHOR.MIDDLE)
    add_paragraph_box(slide, 0.85, 3.50, 8.30, 1.03, [
        {"text": "Important interpretation", "size": 15, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 3},
        {"text": "One positive or negative Δv does not alone identify a conductive or resistive target: its sign depends on the drive/sense geometry. Infer location and sign from all 208 values together.", "size": 12.1, "align": PP_ALIGN.CENTER},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.10)
    add_notes(slide,
        "We use time-difference EIT. v0 is the homogeneous or initial reference frame. v1 is measured after the target is present. The algorithms reconstruct from delta-v. Do not interpret one voltage sign in isolation because electrode geometry changes the sign and sensitivity of each channel.",
        "Đây là time-difference EIT. v0 là trạng thái reference, v1 là trạng thái có target, và đầu vào thuật toán là Δv. Một phần tử Δv âm hoặc dương riêng lẻ chưa đủ kết luận target dẫn tốt hay dẫn kém; phải dùng toàn bộ 208 giá trị cùng sensitivity model.")

    # 22 — controlled simulation.
    slide = add_content_slide(prs, "Why Simulate Before Using CN0565 Hardware?", 22, title_size=25)
    stages = [
        ("1", "Define ground truth", "Known σ at every mesh element"),
        ("2", "Forward solve", "Predict v₀ and v₁ at the boundary"),
        ("3", "Reconstruct", "Give the same Δv to BP, JAC, GREIT"),
        ("4", "Compare", "Check each image against the known answer"),
    ]
    for i, (num, heading, body) in enumerate(stages):
        x = 0.48 + i * 2.37
        add_text(slide, x, 1.05, 2.00, 0.46, num, size=19, bold=True, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON, line=MAROON, rounded=True)
        add_text(slide, x, 1.60, 2.00, 0.45, heading, size=16, bold=True, color=MAROON, align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, x, 2.10, 2.00, 0.92, body, size=13, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=LIGHT, line=MID, rounded=True)
        if i < 3:
            add_arrow(slide, x + 2.03, 2.37, 0.30, 0.22)
    add_paragraph_box(slide, 0.85, 3.48, 8.30, 1.02, [
        {"text": "The simulation provides an answer key.", "size": 17, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 3},
        {"text": "Because the true object is known, output differences can be attributed to the algorithm and its settings—not unknown hardware errors.", "size": 12.6, "align": PP_ALIGN.CENTER},
    ], fill=MAROON_LIGHT, line=MID, rounded=True, margin=0.11)
    add_notes(slide,
        "Simulation is not a replacement for hardware testing. It is a controlled algorithm experiment. Because the conductivity map is known, I can quantify whether an algorithm localizes the target, preserves its sign, resolves two targets, or amplifies noise.",
        "Simulation không thay thế hardware. Nó tạo một bài test có đáp án: ta biết target thật ở đâu nên có thể đánh giá localization, sign, resolution và noise sensitivity của thuật toán.")

    # 23 — model parameters.
    slide = add_content_slide(prs, "How to Read the Conductivity Model", 23)
    add_picture_contain(slide, assets / "ground_truth.png", 0.32, 0.72, 5.30, 4.05)
    add_paragraph_box(slide, 5.72, 0.79, 3.62, 3.94, [
        {"text": "Relative conductivity", "size": 15.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "σrel = σ/σbg. Background = 1 by normalization; it is not automatically water, air, tissue, or 1 S/m.", "size": 11.3, "after": 5},
        {"text": "Target radius = 0.12R", "size": 15.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "R is tank radius. Target radius is 12% of R; target diameter is 24% of R.", "size": 11.3, "after": 5},
        {"text": "Normalized coordinates", "size": 15.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "x/R and y/R run from −1 to +1. (0.5, −0.5) means half a radius right and half a radius down.", "size": 11.3, "after": 5},
        {"text": "Log colour scale", "size": 15.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "log₁₀(σ/σbg): 10× → +1, 1× → 0, 0.1× → −1, so reciprocal contrasts are visually symmetric.", "size": 11.3},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.12)
    add_notes(slide,
        "Background equals one because the simulation is normalized. We divide every conductivity by the background value. The target value ten means ten times the background conductivity, and therefore one tenth of the background resistivity. Radius 0.12R and coordinates are also normalized by the tank radius. The base-ten log maps ten-times and one-tenth to plus one and minus one.",
        "Background = 1 chỉ là chuẩn hóa. perm = 10 nghĩa là conductivity gấp 10 lần background và resistivity bằng 1/10. r = 0.12R nghĩa là bán kính object bằng 12% bán kính tank. Trục x/R, y/R là tọa độ chuẩn hóa. log10 giúp 10× và 0.1× nằm đối xứng +1 và −1.")

    # 24 — mesh and h0.
    slide = add_content_slide(prs, "Where 583 Nodes and 1083 Elements Come From", 24, title_size=24)
    add_picture_contain(slide, assets / "mesh_sweep.png", 0.28, 0.76, 9.30, 2.55)
    add_card(slide, 0.62, 3.48, 2.68, 1.04, "h0 = 0.08", "Requested characteristic triangle spacing in the unit-radius geometry.", accent=MAROON, fill=MAROON_LIGHT, title_size=20, body_size=12.5)
    add_card(slide, 3.63, 3.48, 2.68, 1.04, "583 nodes", "Locations where the FEM solves electric potential u.", accent=NAVY, fill=LIGHT, title_size=20, body_size=12.5)
    add_card(slide, 6.64, 3.48, 2.68, 1.04, "1083 elements", "Triangles carrying piecewise conductivity values σ.", accent=RED, fill=RGBColor(255, 241, 244), title_size=20, body_size=12.5)
    add_source(slide, "Counts are outputs of mesh.create(16, h0=0.08), not universal EIT constants. Smaller h0 increases computation and reduces discretization error.")
    add_notes(slide,
        "The numbers 583 and 1083 are not chosen independently. I specify the circle, sixteen electrode locations, and h0 equals 0.08. The distmesh-style generator then returns a valid triangular mesh. Nodes hold potential unknowns; elements hold conductivity. A smaller h0 creates a finer mesh and more unknowns, but after a point it costs more computation without adding measurement information.",
        "583 và 1083 là output của mesh generator. Input là geometry, 16 electrodes và h0 = 0.08. Node là nơi giải potential; element là tam giác giữ conductivity. Giảm h0 làm mesh mịn hơn, tăng số node/element và thời gian tính, nhưng không tạo thêm dữ liệu đo.")

    # 25 — voltage chart.
    slide = add_content_slide(prs, "How to Interpret the Synthetic Voltage Chart", 25, title_size=24)
    add_picture_contain(slide, assets / "voltage_frames.png", 0.35, 0.95, 6.05, 3.55)
    add_paragraph_box(slide, 6.45, 0.80, 3.00, 3.90, [
        {"text": "Top: v₀ and v₁", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "Each curve has 208 channels: 16 repeated drive groups × 13 sense values.", "size": 10.6, "after": 5},
        {"text": "Why they nearly overlap", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "The target only perturbs the baseline field; Δv is smaller than the absolute model voltage.", "size": 10.6, "after": 5},
        {"text": "Bottom: Δv", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "The algorithms use the full 208-value shape. Peaks mark drive/sense channels that are strongly affected.", "size": 10.6, "after": 5},
        {"text": "Not hardware volts", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "The y-axis is ideal unit-current model output, not yet calibrated to CN0565 amplitude.", "size": 10.6},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.11)
    add_notes(slide,
        "The x-axis is channel index, not time or physical position. Channels zero to twelve belong to the first drive, the next thirteen to the second drive, and so on. The top plot confirms that v1 is a small perturbation of v0. The lower plot isolates that perturbation. A peak means that a particular drive-sense geometry is sensitive to the target, not that the target is located at that x-axis index.",
        "Trục x là index của 208 channels, không phải thời gian hay vị trí target. Mỗi block 13 channels thuộc một drive pair. Hai đường trên gần nhau vì target chỉ perturb baseline. Plot dưới là Δv dùng để reconstruct. Peak chỉ nói channel đó nhạy với target, không phải target nằm tại index đó.")

    # 26 — algorithm overview.
    slide = add_content_slide(prs, "Three Algorithms, One Shared Input", 26)
    add_text(slide, 0.75, 0.74, 8.50, 0.42, "All three receive the same mesh, protocol, v₀, and v₁. None is told the target location.", size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER, margin=0)
    add_card(slide, 0.48, 1.42, 2.85, 2.38, "BP — smear map", "Paint each voltage change back into every region that could have produced it.\n\nFast baseline; broad and blurry.", accent=MAROON, fill=RGBColor(250, 242, 245), title_size=16.5, body_size=12.3)
    add_card(slide, 3.58, 1.42, 2.85, 2.38, "JAC — Jacobian map", "Use the lookup table of how each triangle affects each channel, then solve backward with a stability penalty.\n\nDetail versus noise trade-off.", accent=MAROON, fill=RGBColor(245, 247, 250), title_size=16.5, body_size=12.3)
    add_card(slide, 6.68, 1.42, 2.85, 2.38, "GREIT — grid map", "Use one matrix to convert the 208-value pattern directly into a regular pixel image.\n\nSmooth and fast at run time.", accent=MAROON, fill=RGBColor(241, 247, 252), title_size=16.5, body_size=12.3)
    add_text(slide, 0.90, 4.15, 8.20, 0.50, "Compare location, sign, spread, separation, artifacts, and noise sensitivity - not only which image looks prettier.", size=15, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_notes(slide,
        "This is a fair comparison because all input data are identical. BP is a direct smearing or backprojection. JAC explicitly uses the derivative of each channel with respect to each mesh element. GREIT precomputes a linear reconstruction matrix designed to produce a regular output image with useful performance characteristics.",
        "So sánh công bằng vì input giống nhau. BP smear ngược signal. JAC dùng sensitivity của từng measurement đối với từng element. GREIT precompute một matrix từ 208 measurements sang pixel image. Không nên chỉ chọn ảnh đẹp nhất; phải đo position, spread, artifacts và noise.")

    # 27 — JAC.
    slide = add_content_slide(prs, "JAC: Jacobian Sensitivity and Regularization", 27, title_size=24)
    add_text(slide, 0.55, 0.84, 4.18, 0.54, "Jacobian J: “If element j changes slightly, how does measurement i change?”", size=16, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_table(slide, 0.58, 1.56, 4.10, 1.52,
              ["", "elem 1", "elem 2", "...", "elem 1083"],
              [["meas 1", "J₁₁", "J₁₂", "...", "J₁,₁₀₈₃"], ["meas 2", "J₂₁", "J₂₂", "...", "J₂,₁₀₈₃"], ["...", "...", "...", "...", "..."], ["meas 208", "J₂₀₈,₁", "J₂₀₈,₂", "...", "J₂₀₈,₁₀₈₃"]],
              widths=[0.78, 0.78, 0.78, 0.55, 1.21], font_size=8.5)
    add_text(slide, 0.60, 3.34, 4.10, 0.78, "Linearized forward model\nΔv ≈ J Δσ + noise", size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=LIGHT, line=MID, rounded=True)
    add_paragraph_box(slide, 4.98, 0.84, 4.42, 3.88, [
        {"text": "Why regularization is necessary", "size": 16, "bold": True, "color": MAROON, "after": 3},
        {"text": "1083 unknown element changes > 208 measurements. Multiple images can fit similar data, and noise makes direct inversion unstable.", "size": 11.4, "after": 5},
        {"text": "One-step solution", "size": 15, "bold": True, "color": MAROON, "after": 2},
        {"text": "Δσ̂ = (JᵀJ + λR)⁻¹ JᵀΔv", "size": 16, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 5},
        {"text": "λ controls the trade-off", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "Small λ: more detail and noise. Large λ: more stable, but more blur and lower amplitude.", "size": 11.4, "after": 5},
        {"text": "R and p define the prior", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "Kotre R weights elements by sensitivity. Current p=0.5 and λ=0.01 are starting values to sweep—not universal constants.", "size": 11.2},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.12)
    add_source(slide, "Adler et al. (2007): λ trades resolution against noise attenuation; p=0.5 is a heuristic compromise between boundary and centre noise bias.")
    add_notes(slide,
        "The Jacobian has 208 rows and 1083 columns in this model. Entry J-i-j predicts the change in voltage channel i for a small conductivity change in element j. Because there are more unknown elements than measurements, direct inversion is unstable. Regularization adds a prior penalty. Lambda must be tuned using noise and resolution tests; the paper supports the trade-off, not a universal lambda of 0.01.",
        "J có 208 hàng và 1083 cột. J(i,j) cho biết measurement i nhạy thế nào với conductivity của element j. Vì unknown nhiều hơn measurements nên inverse problem không duy nhất và rất nhạy noise. Regularization thêm penalty để ổn định. Paper chứng minh trade-off của λ, nhưng không chứng minh λ=0.01 luôn đúng; phải sweep.")

    # 28 — BP and GREIT.
    slide = add_content_slide(prs, "BP and GREIT: Two Different Linear Maps", 28, title_size=24)
    add_paragraph_box(slide, 0.55, 0.84, 4.20, 3.88, [
        {"text": "Back Projection (BP)", "size": 18.5, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 4},
        {"text": "Δσ̂ = −HBP Δv", "size": 19, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 5},
        {"text": "HBP is a smear matrix: each measurement contributes to many sensitive nodes.", "size": 12.3, "after": 6},
        {"text": "Current setting", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "weight='none', normalize=True", "size": 12.5, "bold": True, "align": PP_ALIGN.CENTER, "after": 6},
        {"text": "Expected behaviour", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "Very fast baseline, but diffuse sensitivity often creates broad or streaked targets.", "size": 12.3},
    ], fill=RGBColor(250, 242, 245), line=MID, rounded=True, margin=0.14)
    add_paragraph_box(slide, 5.00, 0.84, 4.42, 3.88, [
        {"text": "GREIT in pyEIT", "size": 18.5, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 4},
        {"text": "image pixels = RGREIT Δv", "size": 19, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 5},
        {"text": "Each output pixel is a weighted sum of all 208 voltage changes: pixel k = Σᵢ Rₖᵢ Δvᵢ.", "size": 12.3, "after": 6},
        {"text": "Current settings", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "n=32, s=20, ratio=0.1, p=0.5, λ=0.01", "size": 12.2, "bold": True, "align": PP_ALIGN.CENTER, "after": 6},
        {"text": "What they mean", "size": 14.5, "bold": True, "color": MAROON, "after": 2},
        {"text": "n samples the output grid; s and ratio shape mesh-to-grid weighting. They change appearance, not the 208 independent data values.", "size": 12.0},
    ], fill=RGBColor(241, 247, 252), line=MID, rounded=True, margin=0.14)
    add_source(slide, "GREIT paper: linear reconstruction can be a fast matrix multiplication and should be evaluated using amplitude, position, resolution, shape, ringing, and noise metrics.")
    add_notes(slide,
        "Both methods are linear after their reconstruction matrix has been built. BP uses a direct smear map. GREIT uses a matrix that maps the measurement vector to a regular grid. In the original GREIT framework the matrix is optimized using forward, target, and noise models. The pyEIT implementation here is a sensitivity-based GREIT-style linear map. Increasing n only draws more pixels; it cannot create information that was not present in the 208 channels.",
        "Sau khi setup, cả BP và GREIT đều là matrix multiplication. BP dùng smear matrix. GREIT map 208 values sang grid pixel. n=32 chỉ là sampling của output grid, không tăng resolution vật lý. s và ratio thay đổi spatial weighting/blur nên cũng phải sensitivity test.")

    # 29 — heat map generation.
    slide = add_content_slide(prs, "How the Reconstruction Becomes a Heat Map", 29, title_size=24)
    steps = [
        ("Δv", "208 × 1\nchange vector"),
        ("H or R", "precomputed\nreconstruction matrix"),
        ("Δσ̂", "node, element,\nor grid values"),
        ("Plot", "place values on\nmesh or 32×32 grid"),
        ("Colour", "red = positive\nblue = negative"),
    ]
    for i, (heading, body) in enumerate(steps):
        x = 0.35 + i * 1.92
        add_text(slide, x, 1.34, 1.45, 0.58, heading, size=21, bold=True, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON if i < 3 else NAVY, line=MAROON if i < 3 else NAVY, rounded=True)
        add_text(slide, x, 2.04, 1.45, 1.02, body, size=13.5, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=LIGHT, line=MID, rounded=True)
        if i < 4:
            add_arrow(slide, x + 1.52, 2.34, 0.31, 0.22)
    add_paragraph_box(slide, 0.75, 3.46, 8.50, 1.07, [
        {"text": "The heat map is an estimate, not a camera image.", "size": 16, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 3},
        {"text": "The algorithm outputs numbers; plotting assigns positions and colours. Interpolation or separate normalization can change apparent smoothness and amplitude, so quantitative tests need consistent scaling.", "size": 12.2, "align": PP_ALIGN.CENTER},
    ], fill=MAROON_LIGHT, line=MID, rounded=True, margin=0.11)
    add_notes(slide,
        "The algorithm first returns numbers, not colours. BP returns node values, JAC returns element values which I interpolate to nodes for plotting, and GREIT returns a regular grid. The plotting code assigns red to positive conductivity change and blue to negative change. In the demonstration figure I normalize each algorithm independently, so I compare location and shape, not raw amplitude.",
        "Thuật toán output số. BP output node values, JAC output element values rồi interpolate, GREIT output grid. Heatmap chỉ là cách đặt các số lên vị trí và tô màu. Nếu normalize từng ảnh riêng thì chỉ nên so sánh location, sign, spread, artifacts, không so raw amplitude.")

    # 30 — current result.
    slide = add_content_slide(prs, "Same Four Targets, Three Reconstructions", 30, title_size=24)
    add_picture_contain(slide, assets / "reconstruction_comparison.png", 0.30, 0.72, 9.35, 3.28)
    add_text(slide, 0.68, 4.06, 2.70, 0.48, "BP: broad / diffuse", size=15.5, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_text(slide, 3.65, 4.06, 2.70, 0.48, "JAC: sharper + artifacts", size=15.5, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_text(slide, 6.62, 4.06, 2.70, 0.48, "GREIT: smooth pixel grid", size=15.5, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_source(slide, "Noise = 0. Each reconstruction is normalized independently. This is a conceptual Figure 14 workflow, not a pixel-level reproduction of CN0565.")
    add_notes(slide,
        "All three methods recover the four general locations and signs in this ideal noise-free example. BP spreads the targets broadly. JAC produces concentrated peaks but also mesh-related structure. GREIT is smoother because it maps to a regular grid with target weighting. This single case is not enough to declare a winner.",
        "Cả ba tìm được vị trí và sign cơ bản khi noise = 0. BP blur rộng, JAC peak tập trung nhưng có artifact theo mesh, GREIT smooth trên regular grid. Một case duy nhất chưa đủ xếp hạng thuật toán.")

    # 31 — additional cases.
    slide = add_content_slide(prs, "Additional Cases Reveal Different Failure Modes", 31, title_size=23)
    add_picture_contain(slide, assets / "multi_case_results.png", 0.30, 0.73, 7.08, 4.05)
    add_paragraph_box(slide, 7.48, 0.82, 1.96, 3.87, [
        {"text": "What changes?", "size": 15, "bold": True, "color": MAROON, "after": 5},
        {"text": "Center", "size": 13, "bold": True, "color": NAVY, "after": 1},
        {"text": "Spread and symmetry", "size": 10.5, "after": 5},
        {"text": "Near boundary", "size": 13, "bold": True, "color": NAVY, "after": 1},
        {"text": "Position bias and distortion", "size": 10.5, "after": 5},
        {"text": "Two close targets", "size": 13, "bold": True, "color": NAVY, "after": 1},
        {"text": "Merge or remain separate?", "size": 10.5, "after": 5},
        {"text": "0.1% noise", "size": 13, "bold": True, "color": NAVY, "after": 1},
        {"text": "Artifacts and noise gain", "size": 10.5},
    ], fill=LIGHT, line=MID, rounded=True, margin=0.10)
    add_notes(slide,
        "The center case shows the point-spread shape. The boundary case exposes radial bias and edge distortion. The two-target case tests spatial resolution; in this setup the targets tend to merge. The noisy case tests stability. These rows already show why one visually attractive four-target image is not a complete comparison.",
        "Center case đo point-spread. Boundary case lộ position bias và distortion. Two close targets kiểm tra resolution và thường bị merge. Noise case kiểm tra stability. Vì vậy cần nhiều case có kiểm soát, không thể kết luận từ một ảnh Figure 14.")

    # 32 — fair test matrix and metrics.
    slide = add_content_slide(prs, "A Fair Test Matrix for BP, JAC, and GREIT", 32, title_size=24)
    add_table(slide, 0.42, 0.82, 9.15, 3.45,
        ["Change one variable", "Suggested values", "Measure", "Why it matters"],
        [
            ["Target position", "r/R = 0, 0.3, 0.6, 0.8", "Position error", "Sensitivity is weaker and less uniform toward the centre."],
            ["Target radius", "0.05R, 0.10R, 0.20R", "Resolution / spread", "Tests whether apparent size follows true size."],
            ["Conductivity ratio", "0.1×, 0.5×, 2×, 5×, 10×", "Sign + amplitude response", "Tests linear range and positive/negative contrast."],
            ["Voltage noise", "0, 0.01%, 0.1%, 0.5%, 1%", "Noise figure / artifacts", "Tests regularization and hardware tolerance."],
            ["Target separation", "0.1R to 0.6R", "Two-peak separability", "Tests ability to distinguish nearby objects."],
            ["Mesh / model mismatch", "h0 and boundary shape", "Robustness", "Avoids testing only on the exact inverse model."],
            ["Algorithm parameters", "λ, p, GREIT n/s/ratio", "Metric curves", "Shows why an algorithm behaves differently."],
        ], widths=[1.65, 2.05, 1.75, 3.70], font_size=10.2)
    add_text(slide, 0.72, 4.44, 8.55, 0.38, "Rule: change one variable at a time; reuse the same v₀/v₁ for all algorithms; report metrics before choosing a winner.", size=13.5, bold=True, color=MAROON, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON_LIGHT, line=MID, rounded=True)
    add_notes(slide,
        "The GREIT paper recommends performance measures such as amplitude response, position error, resolution, shape deformation, ringing, and noise amplification. I will use the same principle for all three algorithms. I should also avoid the inverse crime by eventually generating data on a different or finer mesh than the reconstruction mesh.",
        "Test đúng là thay một variable mỗi lần, giữ nguyên data giữa ba thuật toán và đo metric. Có thể dùng position error, resolution/spread, amplitude response, shape deformation, ringing và noise amplification. Sau đó nên tránh inverse crime bằng forward mesh khác reconstruction mesh.")

    # 33 — evidence vs chosen settings.
    slide = add_content_slide(prs, "What the Papers Support - and What We Must Validate", 33, title_size=23)
    add_table(slide, 0.42, 0.84, 9.15, 3.54,
        ["Supported concept", "Evidence", "Current implementation choice"],
        [
            ["Boundary current + boundary voltage measurements", "pyEIT paper, Fig. 2 and Section 2", "16 electrodes; adjacent drive and adjacent sense"],
            ["FEM forward model and Jacobian sensitivity", "pyEIT paper; Adler 2007", "2D unit circle; point-electrode model"],
            ["Regularization trades resolution against noise", "Adler 2007; GREIT 2009", "JAC: method='kotre', p=0.5, λ=0.01"],
            ["Linear reconstruction matrices enable fast imaging", "Adler 2007 and GREIT 2009", "BP/JAC/GREIT one-step difference imaging"],
            ["GREIT should be evaluated by performance metrics", "Adler 2009", "n=32, s=20, ratio=0.1 in pyEIT"],
            ["Mesh density, target geometry, and exact contrasts", "Not fixed by these papers or CN0565 Fig. 14", "h0=0.08; r=0.12R; selected positions and ratios"],
        ], widths=[3.00, 2.65, 3.50], font_size=10.4)
    add_text(slide, 0.78, 4.50, 8.45, 0.32, "A paper can justify the model or trade-off. It usually does not prove one parameter value is universally optimal for this hardware and medium.", size=13.3, bold=True, color=MAROON, align=PP_ALIGN.CENTER, margin=0)
    add_notes(slide,
        "This slide separates evidence from configuration. The papers support the measurement model, FEM, Jacobian, regularization trade-off, linear reconstruction, and evaluation metrics. They do not establish that h0 equals 0.08 or lambda equals 0.01 is optimal for CN0565. Those values are hypotheses or starting settings that must be tested with parameter sweeps and physical data.",
        "Cần phân biệt paper support concept với prove parameter. Paper support EIT model, FEM, Jacobian, regularization trade-off và metrics. Nhưng h0=0.08, λ=0.01, n=32 hay object r=0.12R là lựa chọn implement ban đầu, phải sweep và validate bằng hardware.")

    # 34 — hardware validation.
    slide = add_content_slide(prs, "Next: Validate the CN0565 Measurement Chain", 34, title_size=24)
    hardware_steps = [
        ("1", "Electrical load", "Known resistors / RC network\nCheck channel order, gain, repeatability"),
        ("2", "Homogeneous saline tank", "Stable conductive medium\nMeasure baseline, drift, and noise"),
        ("3", "Known inclusions", "Move conductive and resistive objects\nCompare position and contrast"),
        ("4", "Human study", "Only after safety, ethics, electrodes,\nand performance are established"),
    ]
    for i, (num, heading, body) in enumerate(hardware_steps):
        x = 0.42 + i * 2.39
        add_text(slide, x, 1.02, 2.05, 0.46, f"{num}. {heading}", size=16.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=MAROON, line=MAROON, rounded=True)
        add_text(slide, x, 1.62, 2.05, 1.62, body, size=13.2, color=TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=LIGHT, line=MID, rounded=True)
        if i < 3:
            add_arrow(slide, x + 2.09, 2.29, 0.25, 0.20)
    add_paragraph_box(slide, 0.78, 3.58, 8.46, 0.98, [
        {"text": "A Petri dish or saline tank is the correct next medium test—not yet a human test.", "size": 15.5, "bold": True, "color": MAROON, "align": PP_ALIGN.CENTER, "after": 3},
        {"text": "It exposes electrode, contact, switching, noise, amplitude, and model errors that simulation cannot test.", "size": 12.2, "align": PP_ALIGN.CENTER},
    ], fill=MAROON_LIGHT, line=MID, rounded=True, margin=0.10)
    add_notes(slide,
        "The saline tank is the bridge between simulation and human measurements. Begin with an electrical phantom or known load to verify the switching matrix and acquisition. Then use homogeneous saline to quantify baseline stability and noise. Finally insert known objects at controlled positions. Human testing comes only after safety, ethics, and repeatable performance are established.",
        "Petri dish/saline tank là bước đúng trước human. Trước đó nên test electrical load để xác minh switching và gain. Saline homogeneous đo baseline/noise/drift. Sau đó đặt inclusion biết rõ vị trí và conductivity. Human chỉ sau safety, ethics và performance ổn định.")

    # 35 — references.
    slide = add_content_slide(prs, "References", 35)
    add_paragraph_box(slide, 0.58, 0.88, 8.82, 3.95, [
        {"text": "[1] B. Liu et al., “pyEIT: A Python Based Framework for Electrical Impedance Tomography,” SoftwareX 7 (2018), 304-308.", "size": 15, "bold": True, "after": 14},
        {"text": "[2] A. Adler et al., “Temporal Image Reconstruction in Electrical Impedance Tomography,” Physiological Measurement 28 (2007), S1-S11.", "size": 15, "bold": True, "after": 14},
        {"text": "[3] A. Adler et al., “GREIT: A Unified Approach to 2D Linear EIT Reconstruction of Lung Images,” Physiological Measurement 30 (2009), S35-S55.", "size": 15, "bold": True, "after": 14},
        {"text": "[4] Analog Devices, “CN0565: Electrical Impedance Tomography Measurement System,” Rev. A, Figure 14.", "size": 15, "bold": True, "after": 14},
        {"text": "[5] pyEIT 1.2.4 source used by scripts/eit_sim_playground.py for BP, JAC, GREIT, protocol, FEM, and mesh configuration.", "size": 15, "bold": True},
    ], fill=WHITE, line=None, rounded=False, margin=0.02)
    add_notes(slide,
        "The first three papers are available in the repository's papers folder. The CN0565 circuit note motivates the four-target example, but it does not publish enough simulation detail for a pixel-level reproduction. The exact solver parameter meanings are also checked against the pyEIT implementation used by the script.",
        "Ba paper đầu nằm trong folder papers. CN0565 Figure 14 là nguồn cảm hứng cho four-target case nhưng không công bố đủ parameter để reproduce pixel-by-pixel. Ý nghĩa parameter cụ thể được đối chiếu thêm với source code pyEIT đang dùng.")

    closing = prs.slides.add_slide(prs.slide_layouts[4])
    closing.shapes.title.text = "Thank you"

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Created {output} with {len(prs.slides)} slides")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--assets", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.assets, args.output)


if __name__ == "__main__":
    main()
